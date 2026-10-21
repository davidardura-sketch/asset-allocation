"""Motor independiente de Streamlit. No ejecuta fórmulas del archivo recibido."""
from __future__ import annotations

from dataclasses import asdict, dataclass, field
from datetime import date, datetime, timedelta
from hashlib import sha256
from io import BytesIO
from math import isfinite
from zipfile import BadZipFile, ZipFile
import calendar

import numpy as np
import openpyxl
from openpyxl.utils.datetime import from_excel

IDS = ('EQ_US', 'EQ_EU', 'EQ_JP', 'EQ_APXJ', 'EQ_EM', 'FI_SOV_IG',
       'FI_COV', 'FI_IG', 'FI_HY', 'CASH_EUR', 'ALT_GOLD')
NAMES = ('Estados Unidos', 'Europa', 'Japón', 'Asia Pacífico ex Japón',
         'Mercados emergentes', 'Deuda soberana IG', 'Cédulas y titulizaciones',
         'Renta fija privada IG', 'Renta fija privada HY', 'Liquidez EUR', 'Oro')


class InputError(ValueError):
    """Error de entrada que impide calcular una asignación válida."""


@dataclass(frozen=True)
class Parameters:
    equity: float = .4
    equity_min: float = .3
    equity_max: float = .6
    bonds: float = .4
    cash: float = .15
    gold: float = .05
    equity_regions: tuple = (.294, .6, .032, .016, .058)
    bond_segments: tuple = (.605, .1075, .1625, .125)
    signal_weights: tuple = (.5, .25, .25)
    trend_weights: tuple = (.2, .3, .35, .15)
    trend_scales: tuple = (.05, .1, .15)
    risk_weights: tuple = (.4, .3, .3)
    erp_neutral: float = .025
    valuation_scale: float = .025
    volatility_scale: float = .5
    drawdown_scale: float = .1
    credit_scale: float = .05
    region_tilt: float = .25
    bond_tilt: float = .2
    bond_signal_weights: tuple = (.4, .25, .35)
    bond_fundamental_weights: tuple = (.6, .4)
    bond_fundamental_scales: tuple = (.5, 1.)
    snapshot_lag: int = 3
    day_basis: int = 360
    oas_unit: str = 'percent'

    def validate(self):
        for key, value in asdict(self).items():
            if key == 'oas_unit':
                continue
            vals = value if isinstance(value, (tuple, list)) else [value]
            if any(isinstance(v, bool) or not isinstance(v, (int, float)) or not isfinite(v) for v in vals):
                raise InputError(f'Parámetro inválido: {key}.')
        for key, length in [('equity_regions', 5), ('bond_segments', 4), ('signal_weights', 3),
                            ('trend_weights', 4), ('risk_weights', 3), ('bond_signal_weights', 3),
                            ('bond_fundamental_weights', 2)]:
            vals = getattr(self, key)
            if len(vals) != length or min(vals) < 0 or abs(sum(vals) - 1) > 1e-8:
                raise InputError(f'{key}: los {length} pesos deben ser no negativos y sumar 100%.')
        if len(self.trend_scales) != 3 or len(self.bond_fundamental_scales) != 2:
            raise InputError('Número de escalas incorrecto.')
        if not (0 <= self.equity_min <= self.equity <= self.equity_max < 1):
            raise InputError('Los límites de RV deben cumplir 0 ≤ mínimo ≤ estratégico ≤ máximo < 100%.')
        if min(self.bonds, self.cash, self.gold) < 0 or abs(self.equity+self.bonds+self.cash+self.gold-1) > 1e-8:
            raise InputError('Los pesos estratégicos deben ser no negativos y sumar 100%.')
        scales = (*self.trend_scales, *self.bond_fundamental_scales, self.valuation_scale,
                  self.volatility_scale, self.drawdown_scale, self.credit_scale)
        if min(scales) <= 0 or not (0 <= self.region_tilt < 1 and 0 <= self.bond_tilt < 1):
            raise InputError('Las escalas deben ser positivas y los tilts estar entre 0% y menos de 100%.')
        if int(self.snapshot_lag) != self.snapshot_lag or not 0 <= self.snapshot_lag <= 31:
            raise InputError('El lag debe ser un número entero entre 0 y 31 días.')
        if self.day_basis not in (360, 365) or self.oas_unit not in ('percent', 'bp'):
            raise InputError('Base de días o unidad OAS incorrecta.')

    @property
    def defensive(self):
        return np.array([self.bonds, self.cash, self.gold]) / (1-self.equity)


@dataclass
class Dataset:
    filename: str
    digest: str
    histories: dict
    snapshots: list
    universe: list
    control_date: date | None
    warnings: list[str] = field(default_factory=list)

    @property
    def latest(self):
        return month_end(max(self.histories['rv']))


@dataclass
class Result:
    dataset: Dataset
    params: Parameters
    cutoff: date
    month: date
    dates: list
    raw: np.ndarray
    levels: np.ndarray
    motor: np.ndarray
    signals: dict
    allocation_cells: dict
    regions: list[dict]
    fixed_income: list[dict]
    blocks: list[dict]
    warnings: list[str]
    checks: list[dict]

    @property
    def score(self):
        return self.signals['F30']

    @property
    def position(self):
        s = self.score
        return ('SOBREPONDERAR++' if s >= .6 else 'SOBREPONDERAR' if s >= .2 else
                'NEUTRAL' if s > -.2 else 'INFRAPONDERAR' if s > -.6 else 'INFRAPONDERAR--')


def month_end(d):
    return date(d.year, d.month, calendar.monthrange(d.year, d.month)[1])


def next_month(d):
    return (month_end(d) + timedelta(days=1)).replace(day=1)


def parse_date(v, epoch=None):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool) and 15000 < v < 100000:
        return from_excel(v, epoch=epoch or openpyxl.utils.datetime.WINDOWS_EPOCH).date()
    if isinstance(v, str):
        for fmt in ('%d/%m/%Y', '%Y-%m-%d', '%Y-%m-%d %H:%M:%S', '%d-%m-%Y'):
            try:
                return datetime.strptime(v.strip(), fmt).date()
            except ValueError:
                pass
    return None


def number(v, label, positive=False, nonnegative=False):
    if isinstance(v, bool) or not isinstance(v, (int, float)) or not isfinite(v):
        raise InputError(f'{label}: falta un valor numérico guardado. Refresca Bloomberg y guarda el Excel con sus valores calculados.')
    if positive and v <= 0 or nonnegative and v < 0:
        raise InputError(f'{label}: valor fuera de rango ({v}).')
    return float(v)


def load_excel(content: bytes, filename='Bloomberg.xlsx') -> Dataset:
    if len(content) > 25*1024*1024:
        raise InputError('El archivo supera el límite de 25 MB.')
    try:
        with ZipFile(BytesIO(content)) as z:
            if sum(i.file_size for i in z.infolist()) > 150*1024*1024:
                raise InputError('El contenido descomprimido supera 150 MB.')
        w = openpyxl.load_workbook(BytesIO(content), data_only=True, read_only=True)
    except (BadZipFile, KeyError, OSError, ValueError) as exc:
        if isinstance(exc, InputError):
            raise
        raise InputError('No se puede abrir el archivo. Usa un .xlsx o .xlsm válido, sin contraseña.') from exc
    try:
        required = ['01_Universo', '02_BBG_Snapshot', '03_Hist_RV', '04_Hist_RF', '05_Hist_FX']
        missing = [n for n in required if n not in w.sheetnames]
        if missing:
            raise InputError('Faltan pestañas: ' + ', '.join(missing))
        histories = {}
        for key, name, width in [('rv', '03_Hist_RV', 6), ('rf', '04_Hist_RF', 5), ('fx', '05_Hist_FX', 2)]:
            s = w[name]
            if s.max_row > 10000:
                raise InputError(f'{name}: se admiten hasta 10.000 filas.')
            out = {}
            for i, row in enumerate(s.iter_rows(min_row=10, max_col=width+1, values_only=True), 10):
                dt = parse_date(row[0], w.epoch)
                if dt is None:
                    # Las cabeceras Bloomberg entre filas 10-12 no son observaciones.
                    if any(isinstance(v, (int, float)) for v in row[1:]):
                        raise InputError(f'{name}!A{i}: fecha ausente o no válida en una fila con datos.')
                    continue
                key_month = month_end(dt)
                if key_month in out:
                    raise InputError(f'{name}: hay dos observaciones para {key_month:%m/%Y}. Se requiere una por mes.')
                values = [number(v, f'{name}, fila {i}, columna {j+2}', positive=not (key == 'rf' and j == 4))
                          for j, v in enumerate(row[1:])]
                out[key_month] = (dt, values)
            if not out:
                raise InputError(f'{name}: no contiene históricos calculados.')
            histories[key] = dict(sorted(out.items()))
        snapshots = {}
        for row in w['02_BBG_Snapshot'].iter_rows(min_row=5, max_row=200, max_col=19, values_only=True):
            if row[0] not in IDS:
                continue
            if row[0] in snapshots:
                raise InputError(f'Snapshot: ID duplicado {row[0]}.')
            snapshots[row[0]] = list(row)
        if set(snapshots) != set(IDS):
            raise InputError('Snapshot incompleto. Faltan: ' + ', '.join(set(IDS)-set(snapshots)))
        sn = [snapshots[key] for key in IDS]
        for i, row in enumerate(sn):
            row[8] = parse_date(row[8], w.epoch)
            if row[8] is None:
                raise InputError(f'{row[0]}: LAST_UPDATE_DT no es una fecha válida.')
            number(row[6], f'{row[0]} PX_LAST', positive=i != 9)
            if i < 5:
                for col, label in [(9,'volatilidad'),(10,'PER'),(11,'PER forward'),(12,'P/B')]:
                    number(row[col], f'{row[0]} {label}', positive=True)
                number(row[13], f'{row[0]} dividendo', nonnegative=True)
            elif i < 9:
                number(row[14], f'{row[0]} Yield')
                number(row[15], f'{row[0]} OAS', nonnegative=True)
                number(row[16], f'{row[0]} duración', positive=True)
        u = {r[0]: list(r) for r in w['01_Universo'].iter_rows(min_row=5, max_row=200, max_col=17, values_only=True) if r[0] in IDS}
        if set(u) != set(IDS):
            raise InputError('El universo debe contener los 11 identificadores previstos.')
        for i, key in enumerate(IDS):
            expected = ['USD','EUR','JPY','USD','USD','EUR','EUR','EUR','EUR','EUR','USD'][i]
            if u[key][9] != expected:
                raise InputError(f'{key}: divisa {u[key][9]} no soportada. El modelo v0.4 espera {expected}.')
        control = parse_date(w['00_Control']['B6'].value, w.epoch) if '00_Control' in w.sheetnames else None
        warnings = []
        if any(u[key][8] == 'Price Index' for key in IDS[:5]):
            warnings.append('Japón, Asia ex Japón y/o emergentes usan índices de precio. Sus señales no incluyen dividendos.')
        if any(str(u[key][15]).startswith('REVISAR_BX') for key in IDS[5:9]):
            warnings.append('Persisten campos RF por confirmar en FLDS. Se usan las analíticas numéricas guardadas, como en v0.4.')
        warnings.append('VIX y avisos contrarian de crédito no se actualizan con esta plantilla y no intervienen en la asignación.')
        return Dataset(filename, sha256(content).hexdigest(), histories, sn, [u[k] for k in IDS], control, warnings)
    finally:
        w.close()


def clip(x):
    return float(np.clip(x, -1, 1))


def calculate(data: Dataset, params: Parameters | None = None, cutoff=None, today=None) -> Result:
    p = params or Parameters()
    p.validate()
    cutoff = month_end(cutoff or data.latest)
    today = today or date.today()
    if cutoff >= today.replace(day=1):
        raise InputError('La fecha de corte debe pertenecer a un mes ya cerrado.')
    h = {key: {dt: row for dt, row in hist.items() if dt <= cutoff} for key, hist in data.histories.items()}
    for key, hist in h.items():
        if not hist or max(hist) != cutoff:
            raise InputError(f'Histórico {key.upper()} sin cierre de {cutoff:%m/%Y}. No se reutiliza el mes anterior.')
        dates = sorted(hist)
        for a, b in zip(dates, dates[1:]):
            if a.year*12+a.month+1 != b.year*12+b.month:
                raise InputError(f'Histórico {key.upper()}: falta un mes entre {a:%m/%Y} y {b:%m/%Y}.')
    months = list(h['rv'])
    if len(months) < 37 or len(h['rf']) < 37:
        raise InputError('Se necesitan al menos 37 cierres consecutivos de RV y RF para calcular la volatilidad de 36 meses.')
    if any(m not in h['fx'] for m in months):
        raise InputError('Faltan divisas para uno o más meses del histórico de RV.')
    if min(h['rf']) < min(h['rv']):
        raise InputError('El histórico de RV debe comenzar antes o a la vez que el de RF.')
    for row in data.snapshots:
        lag = (row[8]-cutoff).days
        if abs(lag) > p.snapshot_lag:
            raise InputError(f'{row[0]}: snapshot del {row[8]:%d/%m/%Y} fuera del lag ±{p.snapshot_lag} días del corte. Refresca el snapshot.')
    warnings = list(data.warnings)
    if data.control_date != cutoff:
        old = data.control_date.strftime('%d/%m/%Y') if data.control_date else 'sin fecha'
        warnings.append(f'Control original {old}. Se ha aplicado el cierre {cutoff:%d/%m/%Y} seleccionado, con históricos alineados por mes.')
    after = [r[0] for r in data.snapshots if r[8] > cutoff]
    if after:
        warnings.append('Snapshot posterior al cierre, dentro del lag permitido: '+', '.join(after)+'. No es un corte estricto al cierre.')
    warnings.append('OAS de entrada en '+ ('porcentaje, convertido a puntos básicos multiplicando por 100.' if p.oas_unit == 'percent' else 'puntos básicos, sin conversión.'))
    n = len(months)
    raw = np.full((n, 13), np.nan)
    dates = []
    for i, m in enumerate(months):
        dates.append(h['rv'][m][0])
        raw[i, :6] = h['rv'][m][1]
        if m in h['rf']:
            raw[i, 6:11] = h['rf'][m][1]
        raw[i, 11:] = h['fx'][m][1]
    levels = np.full((n, 15), np.nan)
    levels[:, 0] = [(dt-date(1899,12,30)).days for dt in dates]
    levels[:, 1:6] = raw[:, :5]/raw[0, :5]*100
    for j in (1,4,5):
        levels[:, j] *= raw[0, 11]/raw[:, 11]
    levels[:, 3] *= raw[0, 12]/raw[:, 12]
    start_rf = next(i for i in range(n) if isfinite(raw[i, 6]))
    for j in range(4):
        levels[start_rf:, j+6] = raw[start_rf:, j+6]/raw[start_rf, j+6]*100
    levels[start_rf, 10] = 100
    for i in range(start_rf+1, n):
        factor = 1+raw[i-1, 10]/100*(dates[i]-dates[i-1]).days/p.day_basis
        if factor <= 0:
            raise InputError('La tasa de caja produce un factor de capitalización no positivo.')
        levels[i, 10] = levels[i-1, 10]*factor
    levels[:, 11] = (raw[:, 5]/raw[:, 11])/(raw[0, 5]/raw[0, 11])*100
    levels[:, 12:] = raw[:, [11,12,10]]
    ew, bw = np.array(p.equity_regions), np.array(p.bond_segments)
    motor = np.full((n,18),np.nan)
    motor[:, 0] = levels[:, 0]
    motor[:, 1] = levels[:, 1:6]@ew
    motor[start_rf:, 2] = levels[start_rf:, 6:10]@bw
    motor[start_rf:, 3] = levels[start_rf:, 9]/levels[start_rf:, 6]
    motor[1:, 4] = motor[1:,1]/motor[:-1,1]-1
    motor[start_rf+1:,5] = motor[start_rf+1:,2]/motor[start_rf:-1,2]-1
    motor[start_rf+1:,14:18] = levels[start_rf+1:,6:10]/levels[start_rf:-1,6:10]-1
    for i in range(n):
        if i >= 6:
            motor[i,6] = np.std(motor[i-5:i+1,4],ddof=1)*np.sqrt(12)
        if i >= 36:
            motor[i,7] = np.std(motor[i-35:i+1,4],ddof=1)*np.sqrt(12)
        if i >= 12:
            motor[i,8] = motor[i,1]/max(motor[i-12:i+1,1])-1
        for horizon, col in [(3,9),(6,10),(12,11)]:
            if i >= horizon:
                motor[i,col] = motor[i,1]/motor[i-horizon,1]-1
        if i >= 9:
            motor[i,12] = np.mean(levels[i,1:6] > np.mean(levels[i-9:i+1,1:6],axis=0))
        if i >= start_rf+6:
            motor[i,13] = motor[i,3]/motor[i-6,3]-1
    if motor[-1,7] <= 0:
        raise InputError('La volatilidad de 36 meses es cero. No puede calcularse el régimen de riesgo.')
    s = {}
    for row, col, scale, weight in zip((6,7,8),(9,10,11),p.trend_scales,p.trend_weights):
        s.update({f'B{row}':motor[-1,col],f'C{row}':scale,f'D{row}':clip(motor[-1,col]/scale),f'E{row}':weight})
        s[f'F{row}'] = s[f'D{row}']*weight
    s.update(B9=motor[-1,12], C9=.5, D9=clip(2*motor[-1,12]-1), E9=p.trend_weights[3])
    s['F9'] = s['D9']*s['E9']; s['F10'] = sum(s[f'F{r}'] for r in range(6,10))
    sn = data.snapshots
    s.update(I6=sum(ew[i]*sn[i][11] for i in range(5)), I7=sum(ew[i]*sn[i][12] for i in range(5)),
             I8=sum(ew[i]*sn[i][9] for i in range(5)), B14=sum(ew[i]/sn[i][11] for i in range(5)),
             B15=raw[-1,10]/100, C16=p.erp_neutral)
    s['B16'] = s['B14']-s['B15']; s['D16'] = clip((s['B16']-p.erp_neutral)/p.valuation_scale)
    vals = (motor[-1,6]/motor[-1,7],motor[-1,8],motor[-1,13])
    scores = (clip((1-vals[0])/p.volatility_scale),clip(1+vals[1]/p.drawdown_scale),clip(vals[2]/p.credit_scale))
    for r,v,sc,scale,weight in zip((20,21,22),vals,scores,(p.volatility_scale,p.drawdown_scale,p.credit_scale),p.risk_weights):
        s.update({f'B{r}':v,f'C{r}':scale,f'D{r}':sc,f'E{r}':weight,f'F{r}':sc*weight})
    s['F23'] = sum(s[f'F{r}'] for r in (20,21,22))
    for r,score,weight in zip((27,28,29),(s['F10'],s['D16'],s['F23']),p.signal_weights):
        s.update({f'D{r}':score,f'E{r}':weight,f'F{r}':score*weight})
    s['F30'] = sum(s[f'F{r}'] for r in (27,28,29))
    s['B31'] = p.equity+s['F30']*((p.equity_max-p.equity) if s['F30']>=0 else (p.equity-p.equity_min))
    regions = []
    for i,r in enumerate(range(35,40)):
        ret6,ret12 = levels[-1,i+1]/levels[-7,i+1]-1, levels[-1,i+1]/levels[-13,i+1]-1
        trend = .4*clip(ret6/p.trend_scales[1])+.6*clip(ret12/p.trend_scales[2])
        value = .6*clip((s['I6']/sn[i][11]-1)/.3)+.4*clip((s['I7']/sn[i][12]-1)/.3)
        vol = sn[i][9]/100; risk = clip((s['I8']/100/vol-1)/.5)
        mixed = .45*trend+.4*value+.15*risk; tilt=1+mixed*p.region_tilt
        row = dict(id=IDS[i],name=NAMES[i],strategic=ew[i],return6=ret6,return12=ret12,trend=trend,
                   pe=sn[i][11],pb=sn[i][12],value=value,volatility=vol,risk=risk,score=mixed,tilt=tilt)
        regions.append(row)
    denom = sum(r['strategic']*r['tilt'] for r in regions)
    for i,(r,row) in enumerate(zip(range(35,40),regions)):
        row['within']=row['strategic']*row['tilt']/denom; row['portfolio']=row['within']*s['B31']
        vals = [row[k] for k in ('strategic','return6','return12','trend','pe','pb','value','volatility','risk','score','tilt','within','portfolio')]
        s.update({f'{chr(67+j)}{r}':v for j,v in enumerate(vals)})
    s.update(C40=1., N40=1., O40=s['B31'])
    oas_scale = 100 if p.oas_unit=='percent' else 1
    s['I11'] = sum(bw[i]*sn[i+5][14] for i in range(4))
    s['I12'] = sum(bw[i]*sn[i+5][15]*oas_scale for i in range(4))
    s['I13'] = sum(bw[i]*sn[i+5][16] for i in range(4))
    if s['I11'] == 0 or s['I12'] <= 0:
        raise InputError('Yield estratégico igual a cero u OAS estratégico no positivo: no puede calcularse el fundamental RF.')
    vols = np.std(motor[-6:,14:18],axis=0,ddof=1)*np.sqrt(12)
    if np.any(vols<=0):
        raise InputError('Uno de los segmentos RF tiene volatilidad de seis meses igual a cero.')
    fixed=[]
    for i in range(4):
        ret6,ret12 = levels[-1,i+6]/levels[-7,i+6]-1,levels[-1,i+6]/levels[-13,i+6]-1
        trend=.4*clip(ret6/.05)+.6*clip(ret12/.1)
        risk=clip((np.mean(vols)/vols[i]-1)/.5)
        y, oas, dur=sn[i+5][14],sn[i+5][15]*oas_scale,sn[i+5][16]
        fundamental=p.bond_fundamental_weights[0]*clip((y/s['I11']-1)/p.bond_fundamental_scales[0])+p.bond_fundamental_weights[1]*clip((oas/s['I12']-1)/p.bond_fundamental_scales[1])
        score=float(np.dot([trend,risk,fundamental],p.bond_signal_weights))
        fixed.append(dict(id=IDS[i+5],name=NAMES[i+5],strategic=bw[i],return6=ret6,return12=ret12,
                          volatility=vols[i],trend=trend,risk=risk,yield_pct=y,oas_bp=oas,duration=dur,
                          fundamental=fundamental,score=score,tilt=1+score*p.bond_tilt))
    denom=sum(r['strategic']*r['tilt'] for r in fixed)
    bond_weight=(1-s['B31'])*p.defensive[0]
    for r,row in zip(range(44,48),fixed):
        row['within']=row['strategic']*row['tilt']/denom;row['portfolio']=row['within']*bond_weight
        row['duration_rf']=row['within']*row['duration'];row['duration_portfolio']=row['portfolio']*row['duration']
        vals=[row[k] for k in ('strategic','return6','return12','volatility','trend','risk','yield_pct','oas_bp','duration','fundamental','score','tilt','within','portfolio','duration_rf','duration_portfolio')]
        s.update({f'{chr(67+j)}{r}':v for j,v in enumerate(vals)})
    s.update(C48=1.,O48=1.,P48=bond_weight,Q48=sum(r['duration_rf'] for r in fixed),R48=sum(r['duration_portfolio'] for r in fixed))
    s.update(I14=s['Q48'],I15=s['Q48']-s['I13'])
    blocks=[dict(name=name,strategic=strategic,recommended=recommended,change=recommended-strategic)
            for name,strategic,recommended in [('Liquidez',p.cash,(1-s['B31'])*p.defensive[1]),('Renta fija',p.bonds,bond_weight),
                                               ('Renta variable',p.equity,s['B31']),('Oro',p.gold,(1-s['B31'])*p.defensive[2])]]
    checks=[dict(control='Asignación total',actual=sum(r['recommended'] for r in blocks),expected=1.),
            dict(control='Detalle RV',actual=sum(r['portfolio'] for r in regions),expected=s['B31']),
            dict(control='Detalle RF',actual=sum(r['portfolio'] for r in fixed),expected=bond_weight),
            dict(control='Contribución duración',actual=s['R48'],expected=bond_weight*s['Q48'])]
    for check in checks:
        check['difference']=check['actual']-check['expected']
        check['status']='OK' if abs(check['difference'])<1e-9 else 'ERROR'
    if any(c['status']=='ERROR' for c in checks):
        raise InputError('La asignación no concilia. No se genera una recomendación.')
    # Resultados que alimentan el Excel con cachés numéricas y el PDF.
    a={'B5':cutoff,'D5':s['F30'],'F5':s['B31'],'J5':'OPERATIVO CON LIMITACIONES' if warnings else 'OK',
       'L5':s['Q48'],'N5':'MÁS LARGA' if s['I15']>.25 else 'MÁS CORTA' if s['I15']<-.25 else 'NEUTRAL'}
    for row,b in zip(range(9,13),blocks):
        a.update({f'B{row}':b['strategic'],f'C{row}':b['recommended'],f'D{row}':b['change']})
    a.update(B13=1.,C13=1.,D13=0.,B23=1.,C23=s['B31'],B32=1.,C32=bond_weight,
             E32=sum(r['within']*r['yield_pct'] for r in fixed),F32=sum(r['within']*r['oas_bp'] for r in fixed),
             G32=s['Q48'],H32=s['R48'])
    for row,r in zip(range(18,23),regions):
        a.update({f'B{row}':r['within'],f'C{row}':r['portfolio'],f'D{row}':r['score']})
    for row,r in zip(range(28,32),fixed):
        a.update({f'{col}{row}':r[k] for col,k in zip('BCDEFGH',('within','portfolio','score','yield_pct','oas_bp','duration','duration_portfolio'))})
    for r,sr in zip((37,38,39),(27,28,29)):
        a.update({f'B{r}':s[f'D{sr}'],f'C{r}':s[f'F{sr}']})
    res=Result(data,p,cutoff,next_month(cutoff),dates,raw,levels,motor,s,a,regions,fixed,blocks,warnings,checks)
    a['H5']=res.position
    return res
