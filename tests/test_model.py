from copy import deepcopy
from dataclasses import replace
from datetime import date
from io import BytesIO
import unittest

import numpy as np
import openpyxl
import xlsxwriter

from aa.model import (Parameters,Dataset,InputError,IDS,NAMES,month_end,
                      next_month,parse_date,calculate,load_excel)
from aa.exports import export_excel
from aa.report import export_pdf


def fixture():
    """Mercados sintéticos. El repositorio no contiene datos de Bloomberg."""
    histories={k:{} for k in ('rv','rf','fx')}
    start=date(2020,1,1)
    for i in range(61):
        m=date(start.year+(start.month-1+i)//12,(start.month-1+i)%12+1,1)
        dt=month_end(m)
        values=[100*np.exp(.005*i+.018*np.sin(i*.57+j*.3)) for j in range(6)]
        histories['rv'][dt]=(dt,values)
        histories['rf'][dt]=(dt,[100*np.exp(.001*i+.006*np.sin(i*.61+j)) for j in range(4)]+[2+.4*np.sin(i*.2)])
        histories['fx'][dt]=(dt,[1.1+.04*np.sin(i*.13),130+5*np.sin(i*.17)])
    snapshots=[];universe=[]
    for i,(key,name) in enumerate(zip(IDS,NAMES)):
        currency=['USD','EUR','JPY','USD','USD','EUR','EUR','EUR','EUR','EUR','USD'][i]
        row=[key,'Bloque',name,'Ticker','Ticker','Total Return',100,currency,date(2025,1,31),12+i,18+i,16+i,2+i*.2,2,3.5+i*.1,.4+i*.08,3+i*.2,'OK','']
        snapshots.append(row)
        u=[None]*17;u[0]=key;u[8]='Total Return';u[9]=currency;universe.append(u)
    return Dataset('sintetico.xlsx','synthetic',histories,snapshots,universe,date(2025,1,31),[])


def workbook_bytes(data=None,duplicate=False,missing_sheet=False):
    d=data or fixture();buf=BytesIO();w=xlsxwriter.Workbook(buf,{'in_memory':True})
    datefmt=w.add_format({'num_format':'yyyy-mm-dd'})
    c=w.add_worksheet('00_Control');c.write_datetime('B6',d.control_date,datefmt)
    u=w.add_worksheet('01_Universo')
    for r,row in enumerate(d.universe,4):u.write_row(r,0,row)
    s=w.add_worksheet('02_BBG_Snapshot')
    for r,row in enumerate(d.snapshots,4):
        for col,v in enumerate(row):
            if isinstance(v,date):s.write_datetime(r,col,v,datefmt)
            else:s.write(r,col,v)
    for k,name in [('rv','03_Hist_RV'),('rf','04_Hist_RF'),('fx','05_Hist_FX')]:
        if missing_sheet and k=='fx':continue
        s=w.add_worksheet(name)
        rows=list(d.histories[k].values())
        if duplicate and k=='rv':rows.append(rows[-1])
        for r,(dt,values) in enumerate(rows,9):
            s.write_datetime(r,0,dt,datefmt);s.write_row(r,1,values)
    w.close();return buf.getvalue()


class ModelTests(unittest.TestCase):
    def test_defaults_and_bounds(self):
        p=Parameters();p.validate()
        self.assertEqual(p.equity_regions,(.294,.6,.032,.016,.058))
        self.assertAlmostEqual(sum(p.defensive),1)
        for p in [replace(p,cash=.4),replace(p,equity_max=.2),replace(p,signal_weights=(1,1,1)),replace(p,volatility_scale=0)]:
            with self.assertRaises(InputError):p.validate()

    def test_reconciliation_and_mutation(self):
        d=fixture();r=calculate(d,today=date(2025,2,10))
        self.assertAlmostEqual(sum(x['recommended'] for x in r.blocks),1)
        self.assertTrue(.3<=r.signals['B31']<=.6)
        self.assertAlmostEqual(sum(x['duration_portfolio'] for x in r.fixed_income),r.blocks[1]['recommended']*r.signals['Q48'])
        d.snapshots[1][11]*=2
        changed=calculate(d,today=date(2025,2,10))
        self.assertNotEqual(r.signals['B31'],changed.signals['B31'])

    def test_missing_month_blocks(self):
        d=fixture();del d.histories['rf'][date(2024,9,30)]
        with self.assertRaisesRegex(InputError,'falta un mes'):calculate(d,today=date(2025,2,10))

    def test_stale_and_future_snapshot_blocks(self):
        d=fixture();d.snapshots[1][8]=date(2024,12,31)
        with self.assertRaisesRegex(InputError,'fuera del lag'):calculate(d,today=date(2025,2,10))
        with self.assertRaisesRegex(InputError,'mes ya cerrado'):calculate(fixture(),today=date(2025,1,31))

    def test_price_series_and_lag_are_visible(self):
        d=fixture();d.warnings=['Índice price return'];d.snapshots[1][8]=date(2025,2,1)
        r=calculate(d,today=date(2025,2,10));self.assertTrue(any('posterior' in t for t in r.warnings))

    def test_oas_units_preserve_meaning(self):
        d=fixture();r=calculate(d,today=date(2025,2,10))
        for row in d.snapshots[5:9]:row[15]*=100
        b=calculate(d,replace(Parameters(),oas_unit='bp'),today=date(2025,2,10))
        self.assertAlmostEqual(r.signals['I12'],b.signals['I12'])
        self.assertAlmostEqual(r.fixed_income[3]['within'],b.fixed_income[3]['within'])

    def test_cash_uses_previous_rate_and_actual_days(self):
        d=fixture();r=calculate(d,today=date(2025,2,10));expected=100.
        dates=list(d.histories['rf'])
        for a,b in zip(dates,dates[1:]):expected*=1+d.histories['rf'][a][1][4]/100*(b-a).days/360
        self.assertAlmostEqual(r.levels[-1,10],expected)

    def test_equity_basket_and_risk_six_month_window(self):
        r=calculate(fixture(),today=date(2025,2,10))
        self.assertTrue(np.allclose(r.motor[:,1],r.levels[:,1:6]@np.array(r.params.equity_regions)))
        expected=np.std(r.motor[-6:,14],ddof=1)*np.sqrt(12)
        self.assertAlmostEqual(r.signals['F44'],expected)

    def test_import_readonly_and_duplicates(self):
        d=load_excel(workbook_bytes());self.assertEqual(d.latest,date(2025,1,31))
        with self.assertRaisesRegex(InputError,'dos observaciones'):load_excel(workbook_bytes(duplicate=True))
        with self.assertRaisesRegex(InputError,'Faltan pestañas'):load_excel(workbook_bytes(missing_sheet=True))
        with self.assertRaises(InputError):load_excel(b'not an xlsx')

    def test_export_cached_values_and_formula_endpoints(self):
        r=calculate(fixture(),today=date(2025,2,10));buf=export_excel(r)
        w=openpyxl.load_workbook(BytesIO(buf));v=openpyxl.load_workbook(BytesIO(buf),data_only=True)
        last=len(r.dates)+5
        self.assertIn(f'O{last-5}:O{last}',w['11_Senales']['F44'].value)
        self.assertAlmostEqual(v['12_Asignacion']['C11'].value,r.signals['B31'])
        self.assertAlmostEqual(v['11_Senales']['F44'].value,r.signals['F44'])
        self.assertEqual(w.sheetnames[0],'12_Asignacion')
        self.assertFalse([(s.title,c.coordinate) for s in v for row in s for c in row if c.data_type=='e'])
        pdf=export_pdf(r);self.assertTrue(pdf.startswith(b'%PDF-'))

    def test_zero_and_missing_are_not_equivalent(self):
        d=fixture();d.histories['rf'][max(d.histories['rf'])][1][-1]=0
        calculate(d,today=date(2025,2,10))
        d.snapshots[0][11]=None
        with self.assertRaisesRegex(InputError,'numérico'):load_excel(workbook_bytes(d))


if __name__=='__main__':unittest.main()
